#!/usr/bin/env python3
"""根据最新流量分析模块重新生成汇报 PPT（15 页，原文件不覆盖）。"""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE = RGBColor(0x2E, 0x75, 0xB6)
LIGHT = RGBColor(0xED, 0xF1, 0xF7)
GRAY = RGBColor(0x59, 0x59, 0x59)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def set_font(run, size=18, bold=False, color=RGBColor(0x24, 0x29, 0x32)):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Microsoft YaHei"
    rpr = run._r.get_or_add_rPr()
    ea = rpr.find(qn("a:ea"))
    if ea is None:
        ea = rpr.makeelement(qn("a:ea"), {})
        rpr.append(ea)
    ea.set("typeface", "Microsoft YaHei")


def add_title(slide, title, page_no):
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.05)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    bar.shadow.inherit = False
    tb = slide.shapes.add_textbox(
        Inches(0.55), Inches(0.14), Inches(11.4), Inches(0.78)
    )
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    set_font(run, 28, True, WHITE)

    ptb = slide.shapes.add_textbox(
        Inches(12.15), Inches(0.30), Inches(0.8), Inches(0.5)
    )
    pp = ptb.text_frame.paragraphs[0]
    pp.alignment = PP_ALIGN.RIGHT
    r2 = pp.add_run()
    r2.text = f"{page_no:02d}"
    set_font(r2, 16, True, WHITE)


def add_bullets(slide, bullets, start=1.35):
    tb = slide.shapes.add_textbox(
        Inches(0.65), Inches(start), Inches(12.0), Inches(7.05 - start)
    )
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        if isinstance(item, tuple):
            lead, body = item
            r1 = p.add_run()
            r1.text = lead
            set_font(r1, 20, True, BLUE)
            r2 = p.add_run()
            r2.text = body
            set_font(r2, 20, False)
        else:
            r = p.add_run()
            r.text = "• " + item
            set_font(r, 20, False)
    return tb


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

slides = []

# 1 封面
s = prs.slides.add_slide(blank)
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid()
bg.fill.fore_color.rgb = NAVY
bg.line.fill.background()
bg.shadow.inherit = False
tb = s.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.3), Inches(1.6))
p = tb.text_frame.paragraphs[0]
r = p.add_run()
r.text = "流量分析业务系统"
set_font(r, 44, True, WHITE)
tb2 = s.shapes.add_textbox(Inches(1.0), Inches(3.5), Inches(11.3), Inches(0.8))
p2 = tb2.text_frame.paragraphs[0]
r2 = p2.add_run()
r2.text = "从安全事件发现到版本化智能研判、按日归档的闭环"
set_font(r2, 22, False, LIGHT)
tb3 = s.shapes.add_textbox(Inches(1.0), Inches(4.4), Inches(11.3), Inches(1.0))
p3 = tb3.text_frame.paragraphs[0]
r3 = p3.add_run()
r3.text = "事件聚合收敛 · 版本化自动研判 · PCAP 证据链 · 按日归档 · 自动报告交付"
set_font(r3, 18, False, RGBColor(0xBD, 0xD7, 0xEE))
slides.append("封面")

def new_slide(title, page, bullets):
    s = prs.slides.add_slide(blank)
    add_title(s, title, page)
    add_bullets(s, bullets)

new_slide("告警处置面临的问题", 2, [
    "告警碎片化：多系统各自告警，信息分散，需要手工拼接上下文",
    "依赖专家经验：处置路径取决于个人能力，新人难上手、质量不稳定",
    "过程不可复盘：判断散落在聊天或口头结论，难审计、难复用",
    "报告成本高：事件结束后人工整理材料，汇报慢、格式不统一",
    "历史告警无归档：事件列表全量膨胀，无法快速聚焦当天工作",
])

new_slide("系统核心能力", 3, [
    ("事件聚合建模：", "同源同目标同类型合并，累计发生次数与量化统计"),
    ("版本化自动研判：", "初版快报、收敛终报、手动刷新，报告可追溯"),
    ("按日归档：", "今日工作视图 + 历史归档按日浏览，服务端分页过滤"),
    ("PCAP 证据链：", "单文件下载、聚合全量 ZIP、管理端节点代理"),
    ("报告交付：", "事件分析报告与资产 IP 月度总结，支持 Word 下载"),
    ("独立部署：", "MySQL 全量存储、IAM 剥离、ARM64 离线部署"),
])

new_slide("系统优势", 4, [
    ("更快：", "推送即聚合、自动生成初版分析，缩短告警到研判时间"),
    ("更准：", "先证据后判断，量化统计支撑风险定级"),
    ("可追溯：", "任务、动作、命令、总结全过程留痕，结论可回看"),
    ("可复用：", "报告版本化与月度总结模板沉淀为团队能力资产"),
    ("可运维：", "按日归档、服务端分页、LLM/MySQL 配置健康测试"),
])

new_slide("安全事件处理链路", 5, [
    "1. ta_node 推送：融合采集节点将规则命中与证据一并推送",
    "2. 指纹聚合：源 IP + 目标 IP + 事件类型合并为一条聚合事件",
    "3. 初版分析：自动生成初版快报（analysis_version=1）",
    "4. 静默收敛：30 分钟无新命中自动收敛并生成终版（version=2）",
    "5. 手动刷新：按需重跑最新分析（version>=3）",
    "6. 每日归档：已收敛旧事件按最后活跃日归档",
    "7. 视图交付：今日视图、历史归档按日浏览、全局搜索全量",
])

new_slide("技术亮点", 6, [
    "报告版本化：analysis_version 区分初版 / 终版 / 手动刷新",
    "量化统计：occurrence_count、quant_stats、severity_basis 支撑研判",
    "收敛机制：last_seen_at + aggregation_closed，新命中自动重开收敛",
    "按日归档：archive_date + archive_jobs，列表服务端分页与条件过滤",
    "证据代理：evidence_nodes 映射、路径校验、大小与超时限制",
    "独立运行：MySQL 全量迁移、LLM/MySQL 测试连接、本地认证",
])

new_slide("自动研判与报告版本化", 7, [
    ("初版快报 initial：", "事件创建后自动生成，快速给出初步研判"),
    ("收敛终报 final：", "静默超时后自动生成，occurrence_count 即最终频次"),
    ("手动刷新 manual：", "人工触发重新分析，版本号递增"),
    ("下载口径：", "Word 仅渲染最新一轮报告，不堆叠历史版本"),
])

new_slide("流量聚合与收敛", 8, [
    "聚合键：SHA256（源 IP | 目标 IP | 归一化事件类型）",
    "累计信息：发生次数、首末时间、命中明细、量化统计",
    "收敛窗口：30 分钟无新增命中自动收敛",
    "重新打开：收敛后再次命中会重开收敛并回到今日视图",
])

new_slide("按日归档与工作视图", 9, [
    "归档标记：events.archive_date，NULL 表示未归档（今日视图）",
    "归档条件：仅已收敛且最后活跃早于今日的事件",
    "定时任务：每日 00:10 自动执行，无需人工触发",
    "存量回填：启动时一次性归档历史已收敛事件",
    "查询口径：列表默认今日，历史按日浏览，搜索默认全量",
])

new_slide("PCAP 证据链", 10, [
    "节点字段：evidence_files 透传 name / type / size / path_ref / sha256",
    "单文件下载：按事件与序号经管理端代理节点证据",
    "全量 ZIP：聚合事件全部命中并发拉取，失败附清单",
    "节点解析：evidence_nodes[device_id] 映射节点地址",
    "信息脱敏：报告与界面仅显示文件名，不暴露机器完整路径",
])

new_slide("数据组织与可追溯机制", 11, [
    "events：事件主档，含聚合、收敛、归档、审核、通报编号字段",
    "messages / tasks / actions / commands / executions：过程留痕",
    "summaries：版本化分析结论（初版 / 终版 / 手动刷新）",
    "archive_jobs：每日归档任务与进度",
    "asset_report_summaries / jobs：资产 IP 月度总结与异步任务",
    "traffic_assets / event_maps：资产档案与聚合键映射",
])

new_slide("分析报告生成机制", 12, [
    "报告输入：事件上下文 + 量化统计 + 证据清单 + 分析过程",
    "报告模块：事件概览、关键证据、影响分析、处置建议、附件清单",
    "Word 交付：下载仅渲染最新一轮分析结果",
    "月度总结：按资产 IP 聚合整月报告，异步任务 + 进度查询 + Word",
])

new_slide("独立部署与运维", 13, [
    "存储：全面迁移 MySQL，auto_migrate 幂等建表与补列",
    "配置：LLM / MySQL 参数可视化配置、健康测试、保存回填",
    "认证：IAM 剥离，本地账户独立运行",
    "交付：ARM64 镜像、docker compose、内网离线部署流程",
    "可靠性：配置挂载可写、启动归档回填、收敛后台扫描",
])

new_slide("下一步方向", 14, [
    "修复节点配置在 MariaDB 上的 SQL 兼容问题",
    "提供总览专用统计接口，替代前端分页近似统计",
    "实现分层报告：管理层摘要 / 安全团队报告",
    "增强多节点证据管理与归档日搜索能力",
    "持续优化风险评分与处置优先级",
])

# 15 总结
s = prs.slides.add_slide(blank)
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid()
bg.fill.fore_color.rgb = NAVY
bg.line.fill.background()
bg.shadow.inherit = False
tb = s.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(11.3), Inches(1.2))
p = tb.text_frame.paragraphs[0]
r = p.add_run()
r.text = "让安全分析从“人找线索”变成“系统组织研判”"
set_font(r, 34, True, WHITE)
tb2 = s.shapes.add_textbox(Inches(1.0), Inches(3.0), Inches(11.3), Inches(2.8))
tf = tb2.text_frame
tf.word_wrap = True
for i, line in enumerate([
    "一线安全人员：自动整理上下文、自动生成报告与证据包",
    "安全团队：分析路径标准化、事件过程可复盘、经验可沉淀",
    "管理层：今日视图、按日归档、月度总结支撑跨事件对比与决策",
]):
    para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    para.space_after = Pt(14)
    rr = para.add_run()
    rr.text = line
    set_font(rr, 20, False, LIGHT)

out = "docs/流量分析业务系统汇报_v2.pptx"
prs.save(out)
print("saved", out, "slides:", len(prs.slides._sldIdLst))
